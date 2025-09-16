import os
import uuid
import datetime
from flask import Flask, render_template, request, redirect, url_for, session, jsonify, flash
from werkzeug.utils import secure_filename
import pyodbc
import bcrypt
import pytesseract
from PIL import Image
import fitz
from docx import Document as DocxDocument
import google.generativeai as genai
import re
import openpyxl
import xlrd


# ------------------ App Setup -----------------------
app = Flask(__name__)
UPLOAD_FOLDER = "uploads"
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
app.secret_key = "super-secret-key"  # change this in production!

# Configure Google Gemini API
genai.configure(api_key="AIzaSyAYH1bmWz4hI1s2rgFXvAR5ygtOmoXf4Cs")
model = genai.GenerativeModel("Gemini 2.5-Flash-Lite")

# ------------------ Database Helper -----------------
def get_db():
    """Return a connection to Azure SQL using ODBC Driver 18 and TLS 1.2."""
    server = "tcp:kit300.database.windows.net,1433"
    database = "living_repository"
    username = "living_repository@kit300"
    password = "{4pplec4r7b$77ery}"

    conn_str = (
        "DRIVER={ODBC Driver 18 for SQL Server};"
        f"SERVER={server};"
        f"DATABASE={database};"
        f"UID={username};"
        f"PWD={password};"
        "Encrypt=yes;"
        "TrustServerCertificate=no;"
        "Connection Timeout=30;"
    )

    conn = pyodbc.connect(conn_str)
    return conn

def dict_cursor(cursor):
    """Return rows as dictionaries."""
    columns = [column[0] for column in cursor.description]
    return [dict(zip(columns, row)) for row in cursor.fetchall()]

def get_project_title(project_id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT title FROM Project WHERE project_id=?", (project_id,))
    row = cur.fetchone()
    cur.close()
    conn.close()
    return row[0] if row else "Unknown Project"

# ------------------ Document & Tag Functions ----------------
def generate_tags(text):
    try:
        model = genai.GenerativeModel("gemini-2.5-flash-lite")
        prompt = f"Extract 5-10 concise topic tags for the following document, respond with the tags only, separated by commas:\n\n{text[:3000]}"
        response = model.generate_content(prompt)
        tags = response.text.strip().split(",")
        return [t.strip() for t in tags if t.strip()]
    except Exception as e:
        print("Tagging error:", e)
        return []

def get_or_create_tag(cur, label, category="auto"):
    cur.execute("SELECT tag_id FROM Tag WHERE label = ?", (label,))
    row = cur.fetchone()
    if row:
        return row[0]
    new_id = str(uuid.uuid4())
    cur.execute("INSERT INTO Tag (tag_id, label, category) VALUES (?, ?, ?)", (new_id, label, category))
    return new_id

def extract_word_text(path):
    text = ""
    doc = DocxDocument(path)
    for para in doc.paragraphs:
        text += para.text + "\n"
    return text

def extract_pdf_text(path):
    text = ""
    pdf = fitz.open(path)
    for page in pdf:
        text += page.get_text()
    pdf.close()
    return text

def extract_excel_text(path):
    ext = path.split(".")[-1].lower()
    text = ""
    if ext == "xlsx":
        wb = openpyxl.load_workbook(path, read_only=True)
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                text += " ".join([str(cell) for cell in row if cell is not None]) + "\n"
        wb.close()
    elif ext == "xls":
        wb = xlrd.open_workbook(path)
        for sheet in wb.sheets():
            for row_idx in range(sheet.nrows):
                row = sheet.row_values(row_idx)
                text += " ".join([str(cell) for cell in row if cell]) + "\n"
    return text

def log_event(user_id, project_id, action, object_type=None, object_id=None, object_name=None):
    """
    Insert an event into Audit_Log for tracking user/project activity.

    Args:
        user_id (str): The ID of the user performing the action.
        project_id (str): The project where the action occurred.
        action (str): The action performed ('create','update','delete','login','logout').
        object_type (str, optional): The type of object ('project','document','chat','link','file').
        object_id (str, optional): The ID of the object interacted with.
        object_name (str, optional): A friendly name for the object (e.g., filename).
    """
    conn = get_db()
    cur = conn.cursor()

    log_id = str(uuid.uuid4())
    timestamp = datetime.datetime.now(datetime.timezone.utc)

    cur.execute("""
        INSERT INTO Audit_Log (log_id, user_id, project_id, action, object_type, object_id, object_name, [timestamp])
        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
    """, (log_id, user_id, project_id, action, object_type, object_id, object_name, timestamp))

    conn.commit()
    cur.close()
    conn.close()

# ------------------ Routes -----------------------
@app.route("/")
def index():
    if "user_id" in session:
        return redirect(url_for("dashboard"))
    return render_template("index.html", current_year=datetime.datetime.now().year)

@app.route("/logout")
def logout():
    if "user_id" in session:
        user_id = session["user_id"]
        # Log logout event
        log_event(user_id, None, action="logout")

    session.clear()
    return redirect(url_for("index"))

# ------------------ Authentication -----------------
@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "GET":
        return render_template("login.html")
    
    email = request.form["email"]
    password = request.form["password"]

    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT user_id, name, password_hash FROM [User] WHERE email = ?", (email,))
    row = cur.fetchone()
    cur.close()
    conn.close()

    if not row:
        return "Invalid email or password"

    user_id, name, hashed = row
    if bcrypt.checkpw(password.encode("utf-8"), hashed.encode("utf-8")):
        session["user_id"] = user_id
        session["name"] = name

        # Log login event
        log_event(user_id, None, action="login")

        return redirect(url_for("dashboard"))
    else:
        return "Invalid email or password"

@app.route("/signup", methods=["GET", "POST"])
def signup():
    if request.method == "GET":
        return render_template("signup.html")

    name = request.form.get("username")
    email = request.form.get("email")
    password = request.form.get("password")
    if not name or not email or not password:
        return "Missing field!"

    user_id = str(uuid.uuid4())
    role = "user"
    hashed = bcrypt.hashpw(password.encode("utf-8"), bcrypt.gensalt())

    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT COUNT(*) FROM [User] WHERE email=?", (email,))
    exists = cur.fetchone()[0]
    if exists > 0:
        return "Email already taken"

    cur.execute(
        "INSERT INTO [User] (user_id, name, email, role, password_hash) VALUES (?, ?, ?, ?, ?)",
        (user_id, name, email, role, hashed.decode("utf-8"))
    )
    conn.commit()
    cur.close()
    conn.close()

    # Log signup event
    log_event(user_id, None, action="create", object_type="user", object_id=user_id, object_name=name)

    return "Account created successfully! <a href='/login'>Login here</a>"


@app.route("/dashboard")
def dashboard():
    if "user_id" not in session:
        return redirect(url_for("login"))
    return render_template("dashboard.html", name=session["name"])

# ------------------ Project Routes -----------------
@app.route("/projects")
def projects():
    if "user_id" not in session:
        return redirect(url_for("login"))

    user_id = session["user_id"]
    conn = get_db()
    cur = conn.cursor()
    cur.execute("""
        SELECT p.project_id, p.title, p.description, up.permission
        FROM Project p
        JOIN User_Project up ON p.project_id = up.project_id
        WHERE up.user_id = ?
    """, (user_id,))
    projects = dict_cursor(cur)
    cur.close()
    conn.close()
    return render_template("projects.html", projects=projects)

@app.route("/create_project", methods=["GET", "POST"])
def create_project():
    if "user_id" not in session:
        return redirect(url_for("login"))
    if request.method == "GET":
        return render_template("create_project.html")

    title = request.form["title"]
    description = request.form["description"]
    project_id = str(uuid.uuid4())

    conn = get_db()
    cur = conn.cursor()
    cur.execute("INSERT INTO Project (project_id, title, description, creation_date) VALUES (?, ?, ?, ?)",
                (project_id, title, description, datetime.datetime.now()))
    cur.execute("INSERT INTO User_Project (user_id, project_id, permission, join_date) VALUES (?, ?, ?, ?)",
                (session["user_id"], project_id, "admin", datetime.datetime.now()))
    conn.commit()
    cur.close()
    conn.close()

    # Log project creation
    log_event(
        user_id=session["user_id"],
        project_id=project_id,
        action="create",
        object_type="project",
        object_id=project_id,
        object_name=title
    )

    return redirect(url_for("projects"))


@app.route("/project/<project_id>", methods=["GET", "POST"])
def project_detail(project_id):
    if "user_id" not in session:
        return redirect(url_for("login"))

    user_id = session["user_id"]
    conn = get_db()
    cur = conn.cursor()

    # Handle project deletion
    if request.method == "POST":
        confirm_delete = request.form.get("confirm_delete")
        if confirm_delete == "yes":
            # Optional: check if user has 'admin' permission for this project
            cur.execute("""
                SELECT permission FROM User_Project
                WHERE user_id=? AND project_id=?
            """, (user_id, project_id))
            perm_row = cur.fetchone()
            if not perm_row or perm_row[0] != "admin":
                cur.close()
                conn.close()
                return "You do not have permission to delete this project.", 403

            # Get project title before deletion
            cur.execute("SELECT title FROM Project WHERE project_id=?", (project_id,))
            row = cur.fetchone()
            project_title = row[0] if row else None

            # Delete the project
            cur.execute("DELETE FROM Project WHERE project_id=?", (project_id,))
            conn.commit()
            cur.close()
            conn.close()

            # Log project deletion
            log_event(
                user_id=user_id,
                project_id=project_id,
                action="delete",
                object_type="project",
                object_id=project_id,
                object_name=project_title
            )

            return redirect(url_for("projects"))

    # Fetch project info
    cur.execute("SELECT title, description FROM Project WHERE project_id=?", (project_id,))
    project_row = cur.fetchone()
    if not project_row:
        cur.close()
        conn.close()
        return "Project not found", 404

    columns = [column[0] for column in cur.description]
    project = dict(zip(columns, project_row))

    # Fetch documents
    cur.execute("""
        SELECT d.doc_id, d.title
        FROM Document d
        JOIN Project_Document pd ON d.doc_id = pd.document_id
        WHERE pd.project_id=?
    """, (project_id,))
    documents = dict_cursor(cur)

    cur.close()
    conn.close()

    return render_template(
        "project_detail.html",
        project_id=project_id,
        project_title=project["title"],
        project_description=project["description"],
        documents=documents,
        current_year=datetime.datetime.now().year
    )


@app.route("/project/<project_id>/share", methods=["POST"])
def share_project(project_id):
    if "user_id" not in session:
        return jsonify({"error": "Unauthorized"}), 401

    # Read JSON instead of form
    data = request.get_json() or {}
    permissions = data.get("permissions")
    expiry_str = data.get("expiry")

    if not permissions or not expiry_str:
        return jsonify({"error": "Permissions and expiry date are required"}), 400

    try:
        # Parse expiry date (set to end of day)
        expiry_date = datetime.datetime.strptime(expiry_str, "%Y-%m-%d")
        expiry_date = expiry_date.replace(hour=23, minute=59, second=59, tzinfo=datetime.timezone.utc)
    except Exception as e:
        return jsonify({"error": "Invalid expiry date"}), 400

    link_id = str(uuid.uuid4())
    created_at = datetime.datetime.now(datetime.timezone.utc)

    conn = get_db()
    cur = conn.cursor()
    cur.execute("""
        INSERT INTO Project_Link (link_id, project_id, permissions, expiry_date, created_at)
        VALUES (?, ?, ?, ?, ?)
    """, (link_id, project_id, permissions, expiry_date, created_at))
    conn.commit()
    cur.close()
    conn.close()

    # Log project sharing
    log_event(
        user_id=session["user_id"],
        project_id=project_id,
        action="create",
        object_type="share_link",
        object_id=link_id,
        object_name=f"{permissions} link"
    )

    share_url = url_for("access_shared_project", link_id=link_id, _external=True)

    return jsonify({
        "share_url": share_url,
        "expires": expiry_date.isoformat(),
        "permissions": permissions
    })

@app.route("/share/<link_id>")
def access_shared_project(link_id):
    if "user_id" not in session:
        return redirect(url_for("login"))

    user_id = session["user_id"]
    conn = get_db()
    cur = conn.cursor()

    # Lookup link with permissions
    cur.execute("""
        SELECT project_id, expiry_date, permissions
        FROM Project_Link
        WHERE link_id=?
    """, (link_id,))
    row = cur.fetchone()

    if not row:
        cur.close()
        conn.close()
        return "Invalid or expired link.", 404

    project_id, expiry_date, permissions = row

    # Ensure expiry_date is timezone-aware
    if expiry_date.tzinfo is None:
        expiry_date = expiry_date.replace(tzinfo=datetime.timezone.utc)

    if datetime.datetime.now(datetime.timezone.utc) > expiry_date:
        # Delete expired link
        cur.execute("DELETE FROM Project_Link WHERE link_id=?", (link_id,))
        conn.commit()
        cur.close()
        conn.close()
        return "This link has expired.", 403

    # Check if user already has access
    cur.execute("""
        SELECT 1 FROM User_Project WHERE user_id=? AND project_id=?
    """, (user_id, project_id))
    exists = cur.fetchone()

    if not exists:
        join_date = datetime.datetime.now(datetime.timezone.utc)
        cur.execute("""
            INSERT INTO User_Project (user_id, project_id, permission, join_date)
            VALUES (?, ?, ?, ?)
        """, (user_id, project_id, permissions, join_date))
        conn.commit()

    cur.close()
    conn.close()

    return redirect(url_for("project_detail", project_id=project_id))

@app.route("/project/<project_id>/ingestion")
def ingestion(project_id):
    if "user_id" not in session:
        return redirect(url_for("login"))

    conn = get_db()
    cur = conn.cursor()

    # Get project info
    cur.execute("SELECT title, description FROM Project WHERE project_id=?", (project_id,))
    project_row = cur.fetchone()
    if not project_row:
        cur.close()
        conn.close()
        return "Project not found"

    project = dict(zip([column[0] for column in cur.description], project_row))

    # Get documents with tags
    cur.execute("""
        SELECT d.doc_id,
               d.title,
               d.upload_date,
               u.name AS author,
               d.status,
               STRING_AGG(t.label, ',') AS tags
        FROM Document d
        JOIN Project_Document pd ON d.doc_id = pd.document_id
        LEFT JOIN [User] u ON d.uploader_id = u.user_id
        LEFT JOIN Document_Tag dt ON d.doc_id = dt.document_id
        LEFT JOIN Tag t ON dt.tag_id = t.tag_id
        WHERE pd.project_id = ?
        GROUP BY d.doc_id, d.title, d.upload_date, u.name, d.status
        ORDER BY d.upload_date DESC
    """, (project_id,))
    documents = dict_cursor(cur)

    cur.close()
    conn.close()

    return render_template("ingestion.html",
                           project=project,
                           project_id=project_id,
                           documents=documents)

@app.route("/project/<project_id>/document/<doc_id>/process", methods=["POST"])
def process_document(project_id, doc_id):
    if "user_id" not in session:
        return jsonify({"success": False, "error": "Unauthorized"}), 401

    user_id = session["user_id"]
    conn = get_db()
    cur = conn.cursor()

    try:
        # fetch document
        cur.execute("SELECT title, ocr_text FROM Document WHERE doc_id = ?", (doc_id,))
        row = cur.fetchone()
        if not row or not row[1]:
            return jsonify({"success": False, "error": "Document has no text"}), 400
        doc_title, doc_text = row

        # fetch project
        cur.execute("SELECT title, description FROM Project WHERE project_id = ?", (project_id,))
        row = cur.fetchone()
        if not row:
            return jsonify({"success": False, "error": "Project not found"}), 404
        project_title, project_description = row

        # Gemini AI
        model = genai.GenerativeModel("gemini-2.5-flash-lite")
        prompt = f"""
        Project: {project_title}
        Project Description: {project_description}

        Document Text:
        {doc_text[:6000]}

        Task: Write a concise summary (3–5 sentences) relevant to the project.
        """
        response = model.generate_content(prompt)
        summary = response.text.strip() if response and response.text else None
        if not summary:
            return jsonify({"success": False, "error": "AI did not return summary"}), 500

        # update Project_Document
        cur.execute(
            "UPDATE Project_Document SET contextual_summary = ? WHERE project_id = ? AND document_id = ?",
            (summary, project_id, doc_id)
        )
        # update Document status
        cur.execute(
            "UPDATE Document SET status = ? WHERE doc_id = ?",
            ("Complete", doc_id)
        )
        conn.commit()

        #Log the event
        log_event(
            user_id=user_id,
            project_id=project_id,
            action="update",
            object_type="document",
            object_id=doc_id,
            object_name=doc_title
        )

        return jsonify({"success": True, "summary": summary})

    except Exception as e:
        conn.rollback()
        print("Error processing document:", e)
        return jsonify({"success": False, "error": str(e)}), 500

    finally:
        cur.close()
        conn.close()


@app.route("/project/<project_id>/upload", methods=["POST"])
def project_upload(project_id):
    if "user_id" not in session:
        return jsonify({"error": "Not logged in"}), 401

    file = request.files["file"]
    title = request.form["title"]
    description = request.form.get("description")
    privacy = request.form.get("privacy", "").strip().lower()
    user_id = session["user_id"]

    filename = secure_filename(file.filename)
    path = os.path.join(UPLOAD_FOLDER, filename)
    file.save(path)

    ext = filename.split(".")[-1].lower()
    if ext in ["pdf"]:
        doc_type = "pdf"
        ocr_text = extract_pdf_text(path)
    elif ext in ["doc", "docx"]:
        doc_type = "word"
        ocr_text = extract_word_text(path)
    elif ext in ["xls", "xlsx"]:
        doc_type = "other"
        ocr_text = extract_excel_text(path)
    elif ext in ["jpg", "jpeg", "png", "tiff"]:
        doc_type = "image"
        ocr_text = pytesseract.image_to_string(Image.open(path))
    else:
        doc_type = "other"
        ocr_text = ""

    ocr_text = re.sub(r"\n{2,}", "\n", ocr_text).strip()
    tags = generate_tags(ocr_text or description or title)

    doc_id = str(uuid.uuid4())
    conn = get_db()
    cur = conn.cursor()
    cur.execute("""
        INSERT INTO Document (doc_id, uploader_id, title, description, url, upload_date, type, ocr_text, privacy_level, status)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (doc_id, user_id, title, description, path, datetime.datetime.now(), doc_type, ocr_text, privacy, "pending"))

    cur.execute("INSERT INTO Project_Document (project_id, document_id) VALUES (?, ?)", (project_id, doc_id))

    for tag in tags:
        tag_id = get_or_create_tag(cur, tag)
        cur.execute("INSERT INTO Document_Tag (document_id, tag_id) VALUES (?, ?)", (doc_id, tag_id))

    #log the event
    log_event(
    user_id=session["user_id"],
    project_id=project_id,
    action="create",
    object_type="Document",
    object_id=doc_id,
    object_name=title)

    conn.commit()
    cur.close()
    conn.close()

    return jsonify({"success": True, "doc_id": doc_id, "title": title, "tags": tags})

@app.route("/project/<project_id>/document/<doc_id>", methods=["GET", "POST"])
def document_detail(project_id, doc_id):
    if "user_id" not in session:
        return redirect(url_for("login"))

    user_id = session["user_id"]
    conn = get_db()
    cur = conn.cursor()

    if request.method == "POST":
        if "delete" in request.form:
            # Fetch document info for logging
            cur.execute("SELECT title FROM Document WHERE doc_id=?", (doc_id,))
            doc_row = cur.fetchone()
            doc_title = doc_row[0] if doc_row else None

            # Delete document
            cur.execute("DELETE FROM Document WHERE doc_id=?", (doc_id,))
            conn.commit()

            # Log the deletion
            log_event(
                user_id=user_id,
                project_id=project_id,
                action="delete",
                object_type="document",
                object_id=doc_id,
                object_name=doc_title
            )

            cur.close()
            conn.close()
            return redirect(url_for("project_detail", project_id=project_id))

    # Get document info with uploader
    cur.execute("""
        SELECT d.doc_id, d.title, d.url, d.upload_date, d.type, d.privacy_level, d.ocr_text,
            u.name as uploader_name,
            pd.project_id,
            p.title as project_title
        FROM Document d
        LEFT JOIN [User] u ON d.uploader_id = u.user_id
        LEFT JOIN Project_Document pd ON d.doc_id = pd.document_id
        LEFT JOIN Project p ON pd.project_id = p.project_id
        WHERE d.doc_id=?
    """, (doc_id,))
    row = cur.fetchone()
    document = dict(zip([column[0] for column in cur.description], row)) if row else None

    cur.close()
    conn.close()

    if not document:
        return "Document not found"

    return render_template("document_detail.html",
                           document=document,
                           project_id=project_id)


@app.route("/project/<project_id>/chats")
def project_chats(project_id):
    if "user_id" not in session:
        return redirect(url_for("login"))

    user_id = session["user_id"]
    conn = get_db()
    cur = conn.cursor()

    # Fetch all chats for this project and user
    cur.execute("""
        SELECT chat_id, chat_name, created_at
        FROM Chat
        WHERE project_id=? AND user_id=?
        ORDER BY created_at DESC
    """, (project_id, user_id))
    chats = dict_cursor(cur)

    cur.close()
    conn.close()

    return render_template("project_chats.html",
                           project_id=project_id,
                           chats=chats,
                           current_year=datetime.datetime.now().year,
                           project_title=get_project_title(project_id))

@app.route("/project/<project_id>/chat/<chat_id>")
def load_chat(project_id, chat_id):
    if "user_id" not in session:
        return jsonify({"error": "Unauthorized"}), 401

    user_id = session["user_id"]
    conn = get_db()
    cur = conn.cursor()

    # Ensure chat belongs to user and project
    cur.execute("""
        SELECT chat_id FROM Chat
        WHERE chat_id=? AND project_id=? AND user_id=?
    """, (chat_id, project_id, user_id))
    if not cur.fetchone():
        cur.close()
        conn.close()
        return jsonify([])

    # Fetch messages
    cur.execute("""
        SELECT sender, content, created_at
        FROM ChatContent
        WHERE chat_id=?
        ORDER BY created_at ASC
    """, (chat_id,))
    messages = dict_cursor(cur)

    cur.close()
    conn.close()

    return jsonify(messages)

@app.route("/project/<project_id>/chat/<chat_id>/send", methods=["POST"])
def send_message(project_id, chat_id):
    if "user_id" not in session:
        return jsonify({"error": "Unauthorized"}), 401

    user_id = session["user_id"]
    message = request.json.get("message", "")

    conn = get_db()
    cur = conn.cursor()

    # --- Save user message ---
    content_id_user = str(uuid.uuid4())
    cur.execute("""
        INSERT INTO ChatContent (content_id, chat_id, sender, content, created_at)
        VALUES (?, ?, 'user', ?, ?)
    """, (content_id_user, chat_id, message, datetime.datetime.now(datetime.timezone.utc)))
    conn.commit()

    # --- Build AI prompt from project documents ---

    cur.execute("""
        SELECT 
            d.title, 
            LEFT(d.ocr_text, 1000) AS ocr_text,
            STRING_AGG(t.label, ',') AS tags
        FROM Document d
        LEFT JOIN Document_Tag dt ON d.doc_id = dt.document_id
        LEFT JOIN Tag t ON dt.tag_id = t.tag_id
        JOIN Project_Document pd ON d.doc_id = pd.document_id
        WHERE pd.project_id = ?
        GROUP BY d.doc_id, d.title, d.ocr_text
    """, (project_id,))
    docs = dict_cursor(cur)  # your helper to convert rows to dict

    context_parts = []
    for d in docs:
        tags_str = f" [tags: {d['tags']}]" if d['tags'] else ""
        context_parts.append(f"Document: {d['title']}{tags_str}\n{d.get('ocr_text','')}")

    lvl1_context = "\n\n".join(context_parts)

    # --- Generate AI response ---
    try:

        prompt1 = f"""

        Please return only a list of document titles that are NOT relevant to the user's question seperated by commas and nothing else.

        User Question: {message}

        Project Documents:

        {lvl1_context}

        """

        model = genai.GenerativeModel("gemini-2.5-flash-lite")

        response = model.generate_content(prompt1)

        lvl1_reply = response.text.strip()

  
  

        lvl1_list = [name.strip() for name in lvl1_reply.split(',')]

        deny_set = set(lvl1_list)

    except Exception as e:

        print("Chat error:", e)

        reply = "Sorry, I had trouble generating a response."


    lvl3_info_parts = []

    for d in docs:

        if d['title'] in deny_set:

            continue

        lvl3_info_parts.append(f"Document: {d['title']}\n{d.get('ocr_text','')[:1000]}")

    lvl3_context = "\n\n".join(lvl3_info_parts)

    try:

        prompt = f"""

        You are assisting a user with project documents.

  

        Project Documents:

        {lvl3_context}

  

        User Question: {message}

        """

        model = genai.GenerativeModel("gemini-2.5-flash-lite")

        response = model.generate_content(prompt)

        reply = response.text.strip()

    except Exception as e:

        print("Chat error:", e)

        reply = "Sorry, I had trouble generating a response."

    # --- Save AI response ---
    content_id_ai = str(uuid.uuid4())
    cur.execute("""
        INSERT INTO ChatContent (content_id, chat_id, sender, content, created_at)
        VALUES (?, ?, 'ai', ?, ?)
    """, (content_id_ai, chat_id, reply, datetime.datetime.now(datetime.timezone.utc)))
    conn.commit()

    cur.close()
    conn.close()

    return jsonify({"reply": reply})

@app.route("/project/<project_id>/chat/new", methods=["POST"])
def new_chat(project_id):
    if "user_id" not in session:
        return jsonify({"error": "Unauthorized"}), 401

    user_id = session["user_id"]
    chat_name = request.json.get("chat_name", "").strip()
    if not chat_name:
        return jsonify({"error": "Chat name is required"}), 400

    chat_id = str(uuid.uuid4())
    conn = get_db()
    cur = conn.cursor()
    cur.execute("""
        INSERT INTO Chat (chat_id, project_id, user_id, chat_name, created_at)
        VALUES (?, ?, ?, ?, ?)
    """, (chat_id, project_id, user_id, chat_name, datetime.datetime.utcnow()))
    conn.commit()
    cur.close()
    conn.close()

    return jsonify({"chat_id": chat_id, "chat_name": chat_name})

# DARCIE ADDED THIS FUNCTION
@app.route("/project/<project_id>/document/<doc_id>/tags", methods=["DELETE"])
def delete_document_tag(project_id, doc_id):
    """Delete a tag from a document"""
    if "user_id" not in session:
        return jsonify({"success": False, "error": "Unauthorized"}), 401

    user_id = session["user_id"]
    data = request.get_json()
    if not data or "tag" not in data:
        return jsonify({"success": False, "error": "Tag parameter required"}), 400

    tag_label = data["tag"].strip()
    if not tag_label:
        return jsonify({"success": False, "error": "Tag cannot be empty"}), 400

    conn = get_db()
    cur = conn.cursor()

    try:
        # Verify the document exists and belongs to the project
        cur.execute("""
            SELECT d.doc_id
            FROM Document d
            JOIN Project_Document pd ON d.doc_id = pd.document_id
            WHERE d.doc_id = ? AND pd.project_id = ?
        """, (doc_id, project_id))

        if not cur.fetchone():
            return jsonify({"success": False, "error": "Document not found"}), 404

        # Find the tag ID
        cur.execute("SELECT tag_id FROM Tag WHERE label = ?", (tag_label,))
        tag_row = cur.fetchone()

        if not tag_row:
            return jsonify({"success": False, "error": "Tag not found"}), 404

        tag_id = tag_row[0]

        # Delete the Document_Tag relationship
        cur.execute("""
            DELETE FROM Document_Tag
            WHERE document_id = ? AND tag_id = ?
        """, (doc_id, tag_id))

        rows_affected = cur.rowcount

        if rows_affected == 0:
            return jsonify({"success": False, "error": "Tag was not associated with this document"}), 400

        conn.commit()

        # Log the tag deletion
        log_event(
            user_id=user_id,
            project_id=project_id,
            action="delete",
            object_type="tag",
            object_id=tag_id,
            object_name=tag_label
        )

        return jsonify({
            "success": True,
            "message": f"Tag '{tag_label}' removed successfully"
        })

    except Exception as e:
        conn.rollback()
        print(f"Error deleting tag: {e}")
        return jsonify({"success": False, "error": "Database error occurred"}), 500

    finally:
        cur.close()
        conn.close()

# DARCIE ADDED THIS FUNCTION!!!!!!
@app.route("/project/<project_id>/document/<doc_id>/tags", methods=["POST"])
def add_document_tag(project_id, doc_id):
    """Add a tag to a document"""
    if "user_id" not in session:
        return jsonify({"success": False, "error": "Unauthorized"}), 401

    data = request.get_json()
    if not data or "tag" not in data:
        return jsonify({"success": False, "error": "Tag parameter required"}), 400

    tag_label = data["tag"].strip()
    if not tag_label:
        return jsonify({"success": False, "error": "Tag cannot be empty"}), 400

    conn = get_db()
    cur = conn.cursor()

    try:
        # Verify document exists and belongs to project
        cur.execute("""
            SELECT d.doc_id, d.title
            FROM Document d
            JOIN Project_Document pd ON d.doc_id = pd.document_id
            WHERE d.doc_id = ? AND pd.project_id = ?
        """, (doc_id, project_id))

        document = cur.fetchone()
        if not document:
            return jsonify({"success": False, "error": "Document not found"}), 404

        # Get or create the tag
        tag_id = get_or_create_tag(cur, tag_label, "manual")

        # Check if the tag is already associated with this document
        cur.execute("""
            SELECT COUNT(*)
            FROM Document_Tag
            WHERE document_id = ? AND tag_id = ?
        """, (doc_id, tag_id))

        if cur.fetchone()[0] > 0:
            return jsonify({"success": False, "error": "Tag already exists for this document"}), 400

        # Add the Document_Tag relationship
        cur.execute("""
            INSERT INTO Document_Tag (document_id, tag_id)
            VALUES (?, ?)
        """, (doc_id, tag_id))

        conn.commit()

        # Log the event
        log_event(
            user_id=session["user_id"],
            project_id=project_id,
            action="update",                
            object_type="tag",
            object_id=tag_id,
            object_name=tag_label
        )

        return jsonify({
            "success": True,
            "message": f"Tag '{tag_label}' added successfully",
            "tag": tag_label
        })

    except Exception as e:
        conn.rollback()
        print(f"Error adding tag: {e}")
        return jsonify({"success": False, "error": "Database error occurred"}), 500

    finally:
        cur.close()
        conn.close()

import datetime

@app.route("/project/<project_id>/timeline")
def project_timeline(project_id):
    if "user_id" not in session:
        return redirect(url_for("login"))

    conn = get_db()
    cur = conn.cursor()

    # Fetch project info
    cur.execute("SELECT title FROM Project WHERE project_id=?", (project_id,))
    row = cur.fetchone()
    if not row:
        cur.close()
        conn.close()
        return "Project not found", 404

    project_title = row[0]

    # Fetch audit log entries for this project, ordered newest first
    cur.execute("""
        SELECT al.[timestamp], al.action, al.object_type, al.object_name, u.name as user_name
        FROM Audit_Log al
        LEFT JOIN [User] u ON al.user_id = u.user_id
        WHERE al.project_id=?
        ORDER BY al.[timestamp] DESC
    """, (project_id,))

    log_entries = []
    columns = [column[0] for column in cur.description]

    # Fixed offset for AEST (UTC+10)
    offset = datetime.timedelta(hours=10)

    for row in cur.fetchall():
        entry = dict(zip(columns, row))

        # Append "d" to action
        if entry.get("action"):
            entry["action"] = entry["action"] + "d"

        # Convert timestamp from UTC → AEST (+10h)
        if entry.get("timestamp"):
            entry["timestamp"] = entry["timestamp"] + offset

        log_entries.append(entry)

    cur.close()
    conn.close()

    return render_template(
        "project_timeline.html",
        project_id=project_id,
        project_title=project_title,
        log_entries=log_entries,
        current_year=datetime.datetime.now().year
    )




# ------------------ Main -----------------
if __name__ == "__main__":
    app.run(debug=True)
